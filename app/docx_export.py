"""ייצוא הסכם לקובץ Word מעוצב: A4, RTL מלא, כותרות סעיפים, מבוא, טבלת חתימות ומספרי עמודים.

המבנה מזוהה מהטקסט הנקי לפי תבניות (סעיף, תת-סעיף, הואיל, לפיכך, חתימות), בלי לבקש מהמודל פורמט מיוחד.
"""
import io
import re
from html.parser import HTMLParser

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.draft import clean_output

FONT = "David"  # ponytail: בלי David (Mac, LibreOffice) התוכנה מחליפה בגופן מערכת; להוסיף altName ב-fontTable אם יפריע
AI_NOTE = "טיוטה שנוצרה על ידי AI. יש לבדוק אותה לפני שימוש."
HEADER_NOTE = "טיוטה שנוצרה על ידי AI. יש לבדוק לפני שימוש"
NOTE_COLOR = "9A3412"
GRAY = "808080"
SIGN_LINE = "_" * 24

SECTION = re.compile(r"^\d+\.\s+\S")  # "1. כותרת"
SUBSECTION = re.compile(r"^\d+\.\d+")  # "1.1 ..."
RECITAL = re.compile(r"^(ו?הואיל)(.*)$")
SIGNATURE = re.compile(r"^(שם|תפקיד|חתימה|תאריך)\s*:")
FIELD = re.compile(r"(שם|תפקיד|חתימה|תאריך)\s*:\s*(.*?)\s*(?=(?:שם|תפקיד|חתימה|תאריך)\s*:|$)")
LAWYER = re.compile(r'^אישור\s+(עו"ד|עורך|רו"ח)')
MARKS_ONLY = re.compile(r"^(\[[^\]]+\]\s*)+$")  # "[חותמת] [חתימה]"
HEADING_MAX = 70  # שורה ממוספרת קצרה מזה היא כותרת סעיף, גם בלי נקודתיים

# Word דוחה קובץ שבו ילדי pPr לא בסדר של הסכמה, ולכן כל הוספה עוברת דרך _ppr_add
PPR_ORDER = ("pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
             "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
             "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
             "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
             "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
             "sectPr", "pPrChange")
TBL_AFTER_BIDI = ("w:tblStyleRowBandSize", "w:tblStyleColBandSize", "w:tblW", "w:jc", "w:tblCellSpacing",
                  "w:tblInd", "w:tblBorders", "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook")


def _ppr_add(ppr, el):
    tag = el.tag.split("}")[1]
    ppr.insert_element_before(el, *("w:" + t for t in PPR_ORDER[PPR_ORDER.index(tag) + 1:]))
    return el


def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(f"w:{k}"), v)
    return e


def _cs_size(rpr, size):
    """גודל הגופן לעברית (complex script). בלי זה Word מתעלם מהגודל בטקסט RTL."""
    sz_cs = rpr.find(qn("w:szCs"))
    if sz_cs is None:
        sz_cs = OxmlElement("w:szCs")
        rpr.get_or_add_sz().addnext(sz_cs)
    sz_cs.set(qn("w:val"), str(int(size * 2)))


def _run(p, text, bold=False, size=None, color=None):
    r = p.add_run(text)
    if bold:
        r.font.bold = r.font.cs_bold = True
    if size:
        r.font.size = Pt(size)
        _cs_size(r._r.get_or_add_rPr(), size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return r


def _para(container, text="", bold=False, size=None, color=None, align=None, before=None, after=None,
          keep_next=False, indent=None):
    p = container.add_paragraph()
    if text:
        _run(p, text, bold, size, color)
    pf = p.paragraph_format
    if align is not None:
        pf.alignment = align
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if keep_next:
        pf.keep_with_next = True
    if indent:
        pf.left_indent = Cm(indent)  # בפסקה RTL ‏Word קורא את w:left כתחילת השורה, כלומר מימין
    return p


def _border(p, *sides, sz="6", color="auto", space="4"):
    bdr = _ppr_add(p._p.get_or_add_pPr(), OxmlElement("w:pBdr"))
    for side in sides:  # הסדר בסכמה: top, left, bottom, right
        bdr.append(_el(f"w:{side}", val="single", sz=sz, space=space, color=color))


def _setup(d, ai_note=True):
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(2.5)
    sec._sectPr.find(qn("w:docGrid")).addprevious(OxmlElement("w:bidi"))  # עמוד RTL

    st = d.styles["Normal"]  # כל הפסקאות, כולל header, footer וטבלאות, יורשות ממנו
    st.font.name, st.font.size = FONT, Pt(12)
    rpr = st.element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:cs"), FONT)
    _cs_size(rpr, 12)
    rpr.get_or_add_rtl()
    pf = st.paragraph_format
    pf.line_spacing, pf.space_after, pf.alignment = 1.5, Pt(6), WD_ALIGN_PARAGRAPH.JUSTIFY
    _ppr_add(st.element.get_or_add_pPr(), OxmlElement("w:bidi"))

    if ai_note:  # דוח נתונים הוא לא טיוטה, ולכן בלי הערת AI
        header = sec.header.paragraphs[0]
        header.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(header, HEADER_NOTE, size=9, color=GRAY)

    footer = sec.footer.paragraphs[0]
    footer.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for part in ("עמוד ", "PAGE", " מתוך ", "NUMPAGES"):
        _field(footer, part) if part.isupper() else _run(footer, part, size=9, color=GRAY)


def _field(p, instr):
    """שדה אמיתי של Word (PAGE, NUMPAGES), שמתעדכן לפי הפריסה."""
    for kind in ("begin", "instr", "separate", "value", "end"):
        r = _run(p, "1" if kind == "value" else "", size=9, color=GRAY)
        if kind == "instr":
            t = OxmlElement("w:instrText")
            t.set(qn("xml:space"), "preserve")
            t.text = f" {instr} "
            r._r.append(t)
        elif kind != "value":
            r._r.append(_el("w:fldChar", fldCharType=kind))


def _note_box(d):
    """הערת ה-AI בראש העמוד הראשון: תיבה מוצללת עם גבול."""
    p = _para(d, AI_NOTE, bold=True, size=11, color=NOTE_COLOR, align=WD_ALIGN_PARAGRAPH.CENTER, after=18)
    ppr = p._p.get_or_add_pPr()
    _border(p, "top", "left", "bottom", "right", sz="12", color=NOTE_COLOR, space="6")
    _ppr_add(ppr, _el("w:shd", val="clear", color="auto", fill="FFF4E0"))


def _table(d, rows, cols):
    """טבלה בלי גבולות, שהעמודה הראשונה שלה מימין."""
    t = d.add_table(rows=rows, cols=cols)
    t._tbl.tblPr.insert_element_before(OxmlElement("w:bidiVisual"), *TBL_AFTER_BIDI)
    return t


def _signatures(d, parties):
    """כל צד בעמודה משלו: השדות שלו, וקו חתימה במקום [___]."""
    cols = min(len(parties), 3)
    t = _table(d, -(-len(parties) // cols), cols)
    for i, party in enumerate(parties):
        cell = t.cell(i // cols, i % cols)
        cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
        for label, value in party.items():
            if label == "חתימה" or not value:
                value = SIGN_LINE
            _para(cell, f"{label}: {value}", after=12)
    _para(d)


def _marks(d, line):
    """"[חותמת] [חתימה]" הופך לשורת קווים, וכל שם מתחת לקו שלו."""
    names = re.findall(r"\[([^\]]+)\]", line)
    t = _table(d, 1, len(names))
    for cell, name in zip(t.rows[0].cells, names):
        cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
        _para(cell, SIGN_LINE, align=WD_ALIGN_PARAGRAPH.CENTER, before=24, after=0)
        _para(cell, name, size=10, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER)


def _line(d, line):
    """פסקה אחת לפי התבנית שלה."""
    center = WD_ALIGN_PARAGRAPH.CENTER
    if LAWYER.match(line):  # בלוק נפרד: קו מעל ומרווח
        _border(_para(d, line, bold=True, before=24, keep_next=True), "top", space="12")
    elif line.startswith("לפיכך"):
        _para(d, line, bold=True, align=center, before=12, after=12, keep_next=True)
    elif m := RECITAL.match(line):
        p = _para(d, keep_next=not m.group(2).strip(" :"))
        _run(p, m.group(1), bold=True)
        if m.group(2):
            _run(p, m.group(2), bold=not m.group(2).strip(" :"))
    elif SUBSECTION.match(line):
        _para(d, line, indent=1)
    elif SECTION.match(line) and (line.endswith(":") or len(line) <= HEADING_MAX):
        _para(d, line, bold=True, before=12, after=3, keep_next=True)
    elif line.endswith(":") and len(line) <= HEADING_MAX:  # כותרת בלי מספר, למשל מעל החתימות
        _para(d, line, bold=True, before=12, keep_next=True)
    elif MARKS_ONLY.match(line):
        _marks(d, line)
    else:
        _para(d, line)


def _start(title):
    d = docx.Document()
    _setup(d)
    _note_box(d)
    t = _para(d, title, bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER, after=18)
    _border(t, "bottom", space="6")
    return d


def _save(d):
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def build(title, body):
    title, body = clean_output(title, body)  # גם טקסט שהעובד ערך ידנית בממשק
    d = _start(title)
    _body(d, body.splitlines())
    return _save(d)


def _body(d, lines):
    """שורות טקסט נקי, כל אחת לפי התבנית שלה (סעיף, הואיל, חתימות...)."""
    lines = [l.strip() for l in lines]
    i = 0
    while i < len(lines):
        if not lines[i]:
            i += 1
        elif SIGNATURE.match(lines[i]):  # שורות חתימה רצופות (עם שורות ריקות ביניהן) הופכות לטבלה אחת
            parties = []
            while i < len(lines) and (not lines[i] or SIGNATURE.match(lines[i])):
                for label, value in FIELD.findall(lines[i]):
                    if not parties or label in parties[-1]:  # שדה שחוזר פותח צד חדש
                        parties.append({})
                    parties[-1][label] = value
                i += 1
            _signatures(d, parties)
        else:
            _line(d, lines[i])
            i += 1


BLOCK_TAGS = {"p", "div", "li", "h1", "h2", "h3", "h4", "blockquote", "ol", "ul"}
HEADING_SIZE = {"h1": 16, "h2": 14, "h3": 13, "h4": 12}
FONT_SIZE = {"1": 8, "2": 10, "3": 12, "4": 14, "5": 18, "6": 24, "7": 36}  # <font size> של execCommand
CSS_SIZE = {"x-small": 8, "small": 10, "medium": 12, "large": 14, "x-large": 18, "xx-large": 24, "xxx-large": 36}


def _css_size(v):
    if v in CSS_SIZE:
        return CSS_SIZE[v]
    if m := re.fullmatch(r"([\d.]+)(px|pt)", v):
        return round(float(m[1]) * (0.75 if m[2] == "px" else 1))


class _Html(HTMLParser):
    """ה-HTML מהעורך באתר ← פסקאות: {align, dir, runs: [(טקסט, עיצוב)]}.
    כל תג נשמר במחסנית עם העיצוב והכיוון שהוא מוריש לתוכן שלו."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.stack = [("", {}, {"align": None, "dir": None}, None)]  # (תג, עיצוב, פסקה, מונה רשימה)
        self.blocks, self.runs, self.props = [], [], None

    def flush(self):
        if "".join(t for t, _ in self.runs).strip():
            self.blocks.append({**self.props, "runs": self.runs})
        self.runs, self.props = [], None

    def handle_starttag(self, tag, attrs):
        if tag == "br":
            return self.flush()
        a = dict(attrs)
        style = dict(
            (k.strip().lower(), v.strip().lower())
            for k, _, v in (s.partition(":") for s in (a.get("style") or "").split(";")) if v
        )
        _, fmt, block, _ = self.stack[-1]
        fmt, block = dict(fmt), dict(block)
        if tag in ("b", "strong") or tag in HEADING_SIZE:
            fmt["bold"] = True
        if tag in ("i", "em"):
            fmt["italic"] = True
        if tag == "u":
            fmt["underline"] = True
        if tag in HEADING_SIZE:
            fmt["size"] = HEADING_SIZE[tag]
        if tag == "font" and a.get("size") in FONT_SIZE:
            fmt["size"] = FONT_SIZE[a["size"]]
        if w := style.get("font-weight"):
            fmt["bold"] = w == "bold" or w.isdigit() and int(w) >= 600
        if style.get("font-style") == "italic":
            fmt["italic"] = True
        if "underline" in style.get("text-decoration", "") + style.get("text-decoration-line", ""):
            fmt["underline"] = True
        if "font-size" in style and (size := _css_size(style["font-size"])):
            fmt["size"] = size
        if align := style.get("text-align") or a.get("align"):
            block["align"] = align
        if d := style.get("direction") or a.get("dir"):
            block["dir"] = d if d in ("rtl", "ltr") else block["dir"]
        if tag in BLOCK_TAGS:
            self.flush()
        self.stack.append((tag, fmt, block, 0 if tag in ("ol", "ul") else None))
        if tag == "li":  # מספור או תבליט לפי הרשימה שמעליו
            parent = next((i for i in range(len(self.stack) - 2, 0, -1) if self.stack[i][3] is not None), None)
            prefix = "• "
            if parent is not None:
                t, f, b, n = self.stack[parent]
                self.stack[parent] = (t, f, b, n + 1)
                prefix = f"{n + 1}. " if t == "ol" else "• "
            self.handle_data(prefix)

    def handle_endtag(self, tag):
        if not any(t == tag for t, *_ in self.stack[1:]):
            return  # תג סוגר בלי פותח
        while self.stack[-1][0] != tag:
            self.stack.pop()
        self.stack.pop()
        if tag in BLOCK_TAGS:
            self.flush()

    def handle_data(self, data):
        text = re.sub(r"\s+", " ", data)  # כמו בדפדפן: רווחים ושורות חדשות ב-HTML הם רווח אחד
        if not text or self.stack[-1][0] in ("script", "style") or not self.runs and not text.strip():
            return
        _, fmt, block, _ = self.stack[-1]
        if self.props is None:
            self.props = block
        self.runs.append((text, fmt))


def parse_html(html):
    p = _Html()
    p.feed(html)
    p.close()
    p.flush()
    return p.blocks


def _jc(align, ltr):
    """Word קורא את w:jc="left" בפסקה RTL כתחילת השורה (ימין), כמו ב-ind. לכן ממפים לפי תחילה/סוף, לא לפי צד."""
    if align in ("center", "justify"):
        return WD_ALIGN_PARAGRAPH.CENTER if align == "center" else WD_ALIGN_PARAGRAPH.JUSTIFY
    if align in ("left", "right", "start", "end"):
        start = align == "start" or align == ("left" if ltr else "right")
        return WD_ALIGN_PARAGRAPH.LEFT if start else WD_ALIGN_PARAGRAPH.RIGHT


def _rich(d, block):
    """פסקה שהעובד עיצב בעורך: העיצוב שלו קובע, בלי זיהוי תבניות."""
    ltr = block["dir"] == "ltr"
    p = _para(d, align=_jc(block["align"], ltr))
    if ltr:
        _ppr_add(p._p.get_or_add_pPr(), _el("w:bidi", val="0"))
    runs = block["runs"]
    runs[0] = (runs[0][0].lstrip(), runs[0][1])
    runs[-1] = (runs[-1][0].rstrip(), runs[-1][1])
    for text, f in runs:
        r = _run(p, text, f.get("bold"), f.get("size"))
        if f.get("italic"):
            r.font.italic = r.font.cs_italic = True
        if f.get("underline"):
            r.font.underline = True
        if ltr:
            r.font.rtl = False


def build_html(title, html):
    """ייצוא מהעורך באתר. פסקה בלי עיצוב עוברת את זיהוי התבניות הרגיל (סעיפים, חתימות...), ופסקה מעוצבת נשמרת כמו שהיא."""
    d = _start(clean_output(title, "")[0])
    plain = []
    for block in parse_html(html):
        if block["align"] or block["dir"] or any(f for _, f in block["runs"]):
            _body(d, plain)
            plain = []
            _rich(d, block)
        else:
            plain.append("".join(t for t, _ in block["runs"]))
    _body(d, plain)
    return _save(d)


def build_table(title, header, rows):
    """דוח נתונים כטבלה מסודרת ב-RTL: אותו עיצוב עמוד, בלי הערת "טיוטת AI" (זה לא טיוטה)."""
    d = docx.Document()
    _setup(d, ai_note=False)
    _border(_para(d, title, bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER, after=18), "bottom", space="6")
    t = _table(d, len(rows) + 1, len(header))
    t.style = d.styles["Table Grid"]
    for i, values in enumerate([header, *rows]):
        for j, value in enumerate(values):
            cell = t.cell(i, j)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after, p.paragraph_format.line_spacing = Pt(3), 1.15
            r = _run(p, value, bold=i == 0)
            if j > 0 and i > 0:
                r.font.rtl = False  # מספרים ותאריכים משמאל לימין: "21/10/2026 07:58" ולא "07:58 21/10/2026"
                r.font.name = "Arial"  # בגופן David הנקודתיים בשעה מוצגות עם רווח ("07: 58")
            if i == 0:
                cell._tc.get_or_add_tcPr().append(_el("w:shd", val="clear", color="auto", fill="EDF4FA"))
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
