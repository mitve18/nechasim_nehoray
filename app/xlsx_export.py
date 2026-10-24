"""ייצוא לקובץ Excel מעוצב (גיליון RTL): עמוד "עלויות", וטבלת הסטטוס בעמוד "סנכרון"."""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

TITLE = "עלויות OpenAI API"
BRAND = "0B2A5C"
SOFT = "EBF4FE"
USD = '"$"#,##0.0000'
KINDS = (("draft", "ניסוח טיוטות"), ("ask", "שאלות על המסמכים"), ("sync", "סנכרון ועיבוד מסמכים"))

HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor=BRAND)
TOTAL_FILL = PatternFill("solid", fgColor=SOFT)
LINE = Border(bottom=Side("thin", color="D0D7E2"))


def _header(ws, row, labels):
    for col, label in enumerate(labels, 1):
        c = ws.cell(row, col, label)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 22


def _row(ws, row, values, total=False):
    for col, v in enumerate(values, 1):
        c = ws.cell(row, col, v)
        c.border = LINE
        if isinstance(v, float):
            c.number_format = USD
        if total:
            c.font, c.fill = Font(bold=True, color=BRAND), TOTAL_FILL


def _save(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_table(title, header, rows):
    """טבלה פשוטה עם כותרת, כמו docx_export.build_table."""
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]  # מגבלה של Excel
    ws.sheet_view.rightToLeft = True
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=16, color=BRAND)
    ws["A2"] = f"הופק {datetime.now():%d/%m/%Y %H:%M}"
    ws["A2"].font = Font(color="6B7280")
    _header(ws, 4, header)
    for r, values in enumerate(rows, 5):
        _row(ws, r, values)
    for col, width in zip("AB", (40, 26)):
        ws.column_dimensions[col].width = width
    return _save(wb)


def build_costs(c):
    """c: הפלט של store.costs."""
    wb = Workbook()
    ws = wb.active
    ws.title = "עלויות"
    ws.sheet_view.rightToLeft = True

    ws["A1"] = TITLE
    ws["A1"].font = Font(bold=True, size=16, color=BRAND)
    ws["A2"] = f"{c['days']} הימים האחרונים · הופק {datetime.now():%d/%m/%Y %H:%M}" + (
        f" · נמדד מ-{c['since']}" if c["since"] else "")
    ws["A2"].font = Font(color="6B7280")
    ws["A3"] = "הערכה לפי הטוקנים של כל קריאה ומחירון OpenAI; החשבון ב-OpenAI הוא הקובע."
    ws["A3"].font = Font(italic=True, size=9, color="6B7280")

    r = 5
    _header(ws, r, ["סוג", "קריאות ל-API", "עלות"])
    for key, label in KINDS:
        r += 1
        _row(ws, r, [label, c[key]["calls"], float(c[key]["cost"] or 0)])
    r += 1
    _row(ws, r, ['סה"כ', sum(c[k]["calls"] for k, _ in KINDS), float(c["total"] or 0)], total=True)

    r += 2
    ws.cell(r, 1, "עלות לכל טיוטה").font = Font(bold=True, size=13, color=BRAND)
    r += 1
    _header(ws, r, ["טיוטה", "נוצרה", "בקשות למודל", "עלות"])
    head = r
    for d in c["drafts"]:
        r += 1
        _row(ws, r, [d["title"] or "ללא כותרת", d["created_at"], d["runs"] or 0, float(d["cost"] or 0)])
    if c["unsaved"]:
        r += 1
        _row(ws, r, ["שיחות שלא נשמרו כטיוטה", None, None, float(c["unsaved"])])
    r += 1
    _row(ws, r, ['סה"כ', None, None, float(sum(d["cost"] or 0 for d in c["drafts"]) + c["unsaved"])], total=True)
    if r > head + 1:
        ws.auto_filter.ref = f"A{head}:D{r - 1}"

    for col, width in zip("ABCD", (48, 18, 16, 14)):
        ws.column_dimensions[col].width = width
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    return _save(wb)
