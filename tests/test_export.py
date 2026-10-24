"""בדיקות לניקוי הפלט ולייצוא ל-Word, על פלט אמיתי של ה-Agent. בלי סודות ובלי רשת.

python tests/test_export.py   (או pytest)
"""
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import docx  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from app.docx_export import AI_NOTE, HEADER_NOTE, build, build_html  # noqa: E402
from app.draft import clean_output  # noqa: E402

TITLE = "כתב התחייבות להקמת חניון לילה במערב ראשון לציון"
BODY = """***כתב התחייבות***

**הואיל:**

1. המשרד הביטחון (להלן: "המשרד"), המחזיק במקרקעין הידועים כתוכנית תת''ל 120 לצורך הקמת חניון לילה במערב ראשון לציון, הסכים להקמת הפרויקט על ידי חברת נתיבי איילון בע"מ (להלן: "החברה"), בכפוף לעמידת החברה בתנאים ובדרישות המפורטים בכתב התחייבות זה.

**לפיכך מצהירה ומתחייבת החברה, באמצעות מורשי החתימה מטעמה, כדלקמן:**

1. **ביצוע העבודות:**
   החברה מתחייבת לבצע את כל העבודות הנדרשות להקמת החניון בהתאם להוראות דין, להשיג את כל ההיתרים, התיאומים והאישורים הנדרשים מהרשויות הרלוונטיות, ובכלל זה תיאום עם עיריית ראשון לציון.

2. **תקופת ההתחייבות:**
   תקופת ההתחייבות תחל מדצמבר 2027 ותסתיים בדצמבר 2030;

3. **תקציב:**
   המשרד יעמוד בתקציב כולל של 2 מיליון ש"ח לשנה במשך שלוש שנים.

4. **ביטוח ואחריות:**
   החברה אחראית לוודא קיום ביטוחים מתאימים לכל תקופת ההתחייבות, לרבות ביטוח אחריות מקצועית.

5. **עבודה מול הרשויות:**
   החברה תבצע תיאום מלא עם עיריית ראשון לציון, לרבות קבלת אישורים הנדרשים להקמת החניון.

6. **ביצוע העבודות:**
   העבודות יבוצעו על ידי החברה או קבלנים מטעמה, תוך הקפדה על כללי בטיחות וללא הפרעה לפעילות השוטפת של המשרד.

7. **שינויים והפרות:**
   שום שינוי של תנאי כתב התחייבות זה לא יהיה בתוקף, אלא אם בוצע באישור המשרד ובאישור החברה מראש ובכתב.

8. **סיום הפרויקט:**
   החברה תודיע בכתב למשרד לאחר סיום ההקמה והפיתוח של החניון.

**מורשי חתימה מטעם חברת נתיבי איילון:**

שם: [___] תפקיד: [___] חתימה: [___]

**אישור עו"ד או רו"ח המאשר:**

אני, הח"מ עו"ד/רו"ח [___], מאשר בזאת כי ה"ה החתומים לעיל מורשים לחתום ולהתחייב בשם החברה.

[חותמת] [חתימה]"""


def test_clean_output():
    title, body = clean_output(TITLE, BODY)
    assert "*" not in title + body and "#" not in title + body
    assert title == TITLE
    assert not body.startswith("כתב התחייבות"), body[:40]  # הכותרת לא חוזרת
    assert body.startswith("הואיל:"), body[:40]
    assert "\nהמשרד הביטחון" in body  # המבוא בלי "1."
    assert "\n1. ביצוע העבודות:" in body and "\n8. סיום הפרויקט:" in body  # הסעיפים מתחילים מ-1
    assert "\nהחברה מתחייבת" in body  # ההזחה של שורות ההמשך הוסרה
    assert "[___]" in body  # שדות למילוי נשארים
    assert clean_output("# כותרת", "- פריט\n__חשוב__")[1] == "פריט\nחשוב"
    assert clean_output("", "שם: [___] תפקיד: [_____]")[1] == "שם: [___] תפקיד: [_____]"


def test_build():
    doc = docx.Document(io.BytesIO(build(TITLE, BODY)))
    text = "\n".join(p.text for p in doc.paragraphs)
    cells = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert "*" not in text + cells and "#" not in text + cells
    assert [p.text for p in doc.paragraphs[:2]] == [AI_NOTE, TITLE]
    assert "כתב התחייבות" not in [p.text for p in doc.paragraphs]  # הכותרת הכפולה הוסרה
    sec = doc.sections[0]
    assert round(sec.page_width.cm, 1) == 21.0 and round(sec.left_margin.cm, 1) == 2.5
    assert sec.header.paragraphs[0].text == HEADER_NOTE
    footer = sec.footer.paragraphs[0]._p.xml
    assert "PAGE" in footer and "NUMPAGES" in footer and 'w:fldCharType="begin"' in footer
    assert len(doc.tables) == 2  # חתימות, חותמת + חתימת עו"ד
    assert "שם: [___]" in cells and "חתימה: ____" in cells
    assert all(t._tbl.tblPr.find(qn("w:bidiVisual")) is not None for t in doc.tables)
    heading = next(p for p in doc.paragraphs if p.text == "1. ביצוע העבודות:")
    assert heading.runs[0].bold


def test_build_html_keeps_editor_formatting():
    html = ('<p>1. ביצוע העבודות:</p><p>החברה מתחייבת לבצע.</p>'
            '<p style="text-align: center;"><b>מודגש</b> ו<i>נטוי</i> ו<u>קו</u> <font size="5">גדול</font></p>'
            '<p dir="ltr">Section <b>A</b></p><p align="left">לשמאל</p><p><br></p>'
            '<ol><li>ראשון</li><li>שני</li></ol><p>שם: [___] תפקיד: [___] חתימה: [___]</p>'
            '<script>alert(1)</script><p>&lt;b&gt;לא תג&lt;/b&gt;</p>')
    doc = docx.Document(io.BytesIO(build_html("<b>הסכם</b>", html)))
    paras = {p.text: p for p in doc.paragraphs}
    assert doc.paragraphs[1].text == "<b>הסכם</b>"  # הכותרת היא טקסט, לא HTML
    assert paras["1. ביצוע העבודות:"].runs[0].bold  # פסקה בלי עיצוב: זיהוי התבניות הרגיל
    rich = paras["מודגש ונטוי וקו גדול"]
    assert rich._p.pPr.jc.get(qn("w:val")) == "center"
    runs = {r.text: r for r in rich.runs}
    assert runs["מודגש"].bold and runs["מודגש"].font.cs_bold and runs["נטוי"].italic and runs["קו"].underline
    assert runs["גדול"].font.size.pt == 18 and not runs[" ו"].bold
    ltr = paras["Section A"]
    assert ltr._p.pPr.find(qn("w:bidi")).get(qn("w:val")) == "0" and ltr.runs[0].font.rtl is False
    assert paras["לשמאל"]._p.pPr.jc.get(qn("w:val")) == "right"  # בפסקה RTL ‏Word קורא right כסוף השורה
    assert "1. ראשון" in paras and "2. שני" in paras
    assert "<b>לא תג</b>" in paras and not any("alert" in p.text for p in doc.paragraphs)
    assert len(doc.tables) == 1 and "חתימה: ____" in doc.tables[0].cell(0, 0).text  # שורת חתימות עדיין טבלה


def test_build_costs_xlsx():
    import openpyxl
    from app.xlsx_export import build_costs
    c = {"days": 7, "since": "01/10/2026", "total": 0.5, "unsaved": 0.1,
         "draft": {"cost": 0.3, "calls": 4}, "ask": {"cost": 0.15, "calls": 2}, "sync": {"cost": 0.05, "calls": 1},
         "drafts": [{"title": "הסכם", "created_at": "20/10/2026 10:00", "runs": 2, "cost": 0.2}]}
    ws = openpyxl.load_workbook(io.BytesIO(build_costs(c))).active
    assert ws.sheet_view.rightToLeft
    rows = [[v for v in r if v is not None] for r in ws.iter_rows(values_only=True)]
    assert ['סה"כ', 7, 0.5] in rows and ["הסכם", "20/10/2026 10:00", 2, 0.2] in rows
    assert rows[-1][0] == 'סה"כ' and abs(rows[-1][1] - 0.3) < 1e-9


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    test_clean_output()
    test_build()
    test_build_costs_xlsx()
    test_build_html_keeps_editor_formatting()
    print("ok")
