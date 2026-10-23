"""בדיקות לדוח "סטטוס מאגר ההסכמים": השאילתה, פורמט הטקסט וה-TSV, וקובץ ה-Word. בלי סודות ובלי רשת.

python tests/test_status_report.py   (או pytest)
"""
import contextlib
import io
import re
import sys
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import docx  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from pg import postgres  # noqa: E402

import sync  # noqa: E402
from app import docx_export, main, store  # noqa: E402

VEC = "[" + ",".join(["0.1"] * 1536) + "]"
ROWS = [{"label": 'סה"כ מסמכים בתיקייה', "value": "48", "tone": ""},
        {"label": "עמודים", "value": "1,236", "tone": ""},
        {"label": "הסכמים שעברו OCR", "value": "—", "tone": "muted"},
        {"label": "אינדוקס אחרון", "value": "21/10/2026 07:58", "tone": ""}]


def skip(reason):
    try:
        import pytest
        pytest.skip(reason)
    except ImportError:
        print("SKIP:", reason)


def _doc(conn, file_id, mime, method, pages, status="indexed", text="[עמוד 1]\n\nt"):
    conn.execute("""insert into agreements (file_id, name, modified_time, method, summary, text, embedding, pages)
                    values (%s, %s, '1', %s, 's', %s, %s::vector, %s)""", (file_id, file_id, method, text, VEC, pages))
    conn.execute("insert into files (file_id, name, mime_type, status) values (%s, %s, %s, %s)",
                 (file_id, file_id, mime, status))


def _values(conn):
    return {r["label"]: (r["value"], r["tone"]) for r in store.status_rows(conn)}


def test_status_rows_query():
    with postgres() as conn:
        if conn is None:
            return skip("אין py-pglite או Node")
        store.init(conn)
        before = _values(conn)  # לפני הסנכרון הראשון: לא ידוע = "—", לא 0
        assert before == {'סה"כ מסמכים בתיקייה': ("—", "muted"), "מאונדקסים": ("0", ""), "נכשלו": ("—", "muted"),
                          "עמודים": ("—", "muted"), "הסכמים שעברו OCR": ("0", ""),
                          "אינדוקס אחרון": ("—", "muted")}, before

        _doc(conn, "pdf", "application/pdf", "text", 1200)
        _doc(conn, "scan", "image/png", "ocr", 2)
        _doc(conn, "word", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx", None)
        _doc(conn, "old-scan", "application/pdf", "ocr", None)  # אונדקס לפני שהעמודים נשמרו
        conn.execute("insert into files (file_id, name, mime_type, status, error) "
                     "values ('bad', 'bad', 'image/png', 'failed', 'OCR נכשל')")
        v = _values(conn)
        assert v['סה"כ מסמכים בתיקייה'] == ("5", "") and v["מאונדקסים"] == ("4", "") and v["נכשלו"] == ("1", "bad")
        assert v["עמודים"] == ("1,202 (חלקי)", "muted"), v  # old-scan עוד לא נספר: הסכום של מה שכבר נספר
        assert v["הסכמים שעברו OCR"] == ("2", "")  # הסכמים ולא עמודים, ולא תלוי בספירת העמודים

        assert store.missing_pages(conn) == {"word", "old-scan"}
        store.save_pages(conn, "old-scan", 5)  # מה שהסנכרון הבא עושה. DOCX נשאר בלי עמודים ולא חוסם
        v = _values(conn)
        assert v["עמודים"] == ("1,207", "") and v["הסכמים שעברו OCR"] == ("2", ""), v

        conn.execute("insert into sync_runs (started_at, finished_at, added, updated, removed, ocr, failed) "
                     "values (now(), now(), 0, 0, 0, 0, 0)")
        assert re.fullmatch(r"\d\d/\d\d/\d{4} \d\d:\d\d", _values(conn)["אינדוקס אחרון"][0])


def test_fill_missing_pages_and_ocr_quality_for_existing_docs():
    with postgres() as conn:
        if conn is None:
            return skip("אין py-pglite או Node")
        store.init(conn)
        _doc(conn, "pdf", "application/pdf", "text", None)
        _doc(conn, "done", "application/pdf", "text", 4)
        _doc(conn, "word", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx", None)
        _doc(conn, "broken", "application/pdf", "text", None)
        _doc(conn, "scan", "image/png", "ocr", 2)  # סרוק בלי ציון: מקבל ציון, העמודים כבר ידועים
        _doc(conn, "scored", "image/png", "ocr", 1)
        store.save_quality(conn, "scored", 90)  # כבר יש ציון: לא מריצים OCR שוב
        listing = [{"id": i, "name": i, "mimeType": m} for i, m in
                   (("pdf", "application/pdf"), ("done", "application/pdf"), ("word", "application/msword"),
                    ("broken", "application/pdf"), ("scan", "image/png"), ("scored", "image/png"))]
        downloads, ocr_runs = [], []

        def fake_count(f, data):
            if f["id"] == "broken":
                raise ValueError("EOF marker not found")
            return 9

        def fake_extract(f, data):
            ocr_runs.append(f["id"])
            return "טקסט", "ocr", 73

        stats = {"pages_counted": 0, "ocr_scored": 0, "fill_failed": []}
        with mock.patch.object(sync.drive, "download", lambda svc, f: downloads.append(f["id"]) or b"x"), \
                mock.patch.object(sync.extract, "page_count", fake_count), \
                mock.patch.object(sync.extract, "extract", fake_extract):
            sync.fill_missing(conn, None, listing, stats)
        assert downloads == ["pdf", "broken", "scan"]  # מה שכבר הושלם, ו-DOCX, לא מורדים
        assert ocr_runs == ["scan"]
        rows = {r[0]: r[1:] for r in conn.execute("select file_id, pages, ocr_quality from agreements").fetchall()}
        assert rows == {"pdf": (9, None), "done": (4, None), "word": (None, None), "broken": (None, None),
                        "scan": (2, 73), "scored": (1, 90)}, rows  # broken יושלם בריצה הבאה
        assert stats == {"pages_counted": 1, "ocr_scored": 1, "fill_failed": ["broken: EOF marker not found"]}

        md = sync.summary({"added": 0, "updated": 0, "removed": 0, "unsupported": 0, "errors": [], "in_db": 6,
                           "folder": "f", "seen": [], "files": {"added": [], "updated": [], "backfilled": []}} | stats)
        assert ("**נספרו עמודים ל-1 מסמכים · חושב ציון איכות OCR ל-1 מסמכים**\n"
                "לא הושלמו (1, ינוסו שוב בריצה הבאה):\n- broken: EOF marker not found") in md, md


def test_fill_page_marks_for_existing_docs():
    with postgres() as conn:
        if conn is None:
            return skip("אין py-pglite או Node")
        store.init(conn)
        _doc(conn, "old", "application/pdf", "text", 3, text="טקסט מלפני סימוני העמודים")
        _doc(conn, "one", "application/pdf", "text", 1, text="עמוד יחיד, בלי סימון")
        _doc(conn, "word", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx", None,
             text="Word בלי עמודים")
        store.save_quality(conn, "old", 0)  # כדי שרק סימוני העמודים יחסרו
        listing = [{"id": i, "name": i, "mimeType": m} for i, m in conn.execute(
            "select file_id, mime_type from files").fetchall()]
        stats = {"pages_counted": 0, "ocr_scored": 0, "pages_marked": 0, "fill_failed": []}
        with mock.patch.object(sync.drive, "download", lambda svc, f: b"x"), \
                mock.patch.object(sync.extract, "extract", lambda f, data: ("[עמוד 1]\n\nא\n\n[עמוד 2]\n\nב", "text", None)):
            sync.fill_missing(conn, None, listing, stats)
        texts = dict(conn.execute("select file_id, text from agreements").fetchall())
        assert texts["old"].startswith("[עמוד 1]") and stats["pages_marked"] == 1
        assert texts["one"] == "עמוד יחיד, בלי סימון" and texts["word"] == "Word בלי עמודים"
        downloads = []  # הריצה הבאה לא מורידה שוב: Word לא נספר (countable), והשאר כבר מסומנים או בעמוד אחד
        with mock.patch.object(sync.drive, "download", lambda svc, f: downloads.append(f["id"]) or b"x"):
            sync.fill_missing(conn, None, listing, stats)
        assert downloads == []


def test_tsv_format():
    tsv = main.status_tsv(ROWS)
    assert tsv.split("\n") == ["מדד\tערך", 'סה"כ מסמכים בתיקייה\t48', "עמודים\t1,236", "הסכמים שעברו OCR\t—",
                               "אינדוקס אחרון\t21/10/2026 07:58"]
    assert all(line.count("\t") == 1 for line in tsv.split("\n"))  # שתי עמודות בכל שורה באקסל


def test_api_status_and_word_export():
    with mock.patch.multiple(store, connect=lambda: contextlib.nullcontext(),
                             status=lambda conn: {"updated_at": "21/10/2026 07:58", "documents": 45, "failed": False},
                             status_rows=lambda conn: ROWS):
        c = TestClient(main.app)
        s = c.get("/api/status").json()
        r = c.get("/api/status/docx")
    assert s["updated_at"] == "21/10/2026 07:58" and s["documents"] == 45  # שורת הסטטוס בראש האתר לא השתנתה
    assert s["title"] == "סטטוס מאגר ההסכמים" and s["rows"] == ROWS and s["tsv"] == main.status_tsv(ROWS)

    assert r.status_code == 200 and r.content[:2] == b"PK"
    assert "filename*=UTF-8''" in r.headers["content-disposition"]
    d = docx.Document(io.BytesIO(r.content))  # נפתח בלי שגיאות
    assert d.sections[0]._sectPr.find(qn("w:bidi")) is not None  # עמוד RTL
    assert d.styles["Normal"].element.pPr.find(qn("w:bidi")) is not None  # פסקאות RTL
    table = d.tables[0]
    assert table._tbl.tblPr.find(qn("w:bidiVisual")) is not None and table.style.name == "Table Grid"
    assert [[c.text for c in row.cells] for row in table.rows] == [["מדד", "ערך"], *([x["label"], x["value"]] for x in ROWS)]
    value_run = table.rows[4].cells[1].paragraphs[0].runs[0]
    assert value_run.font.rtl is False  # "21/10/2026 07:58" משמאל לימין
    everything = "\n".join([p.text for p in d.paragraphs] + [p.text for p in d.sections[0].header.paragraphs])
    assert docx_export.AI_NOTE not in everything and "AI" not in everything  # דוח נתונים, לא טיוטה
    assert d.paragraphs[0].text == "סטטוס מאגר ההסכמים"


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    for test in [v for k, v in dict(globals()).items() if k.startswith("test_")]:
        test()
        print("ok ", test.__name__)
