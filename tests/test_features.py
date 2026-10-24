"""בדיקות ליכולות החדשות: טבלת files בסנכרון, backfill, דוחות, העלאה, מסמך בסיס, הגבלת הסנכרון ו-OCR לתמונות.

בדיקות ה-SQL רצות מול Postgres אמיתי (PGlite, ראה tests/pg.py). Drive, OpenAI ו-GitHub מוחלפים בתחליפים.
בלי סודות ובלי רשת.

python tests/test_features.py   (או pytest)
"""
import contextlib
import hashlib
import io
import os
import shutil
import sys
import threading
import time
import zipfile
from datetime import date
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from agno.agent import Agent  # noqa: E402
from agno.run.base import RunStatus  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from pg import postgres  # noqa: E402

import sync  # noqa: E402
from app import draft, drive, extract, github, main, store  # noqa: E402

PDF, PNG = "application/pdf", "image/png"
VEC = "[" + ",".join(["0.1"] * 1536) + "]"
META = {"doc_type": "הסכם", "parties": ["א", "ב"], "plan_number": None, "location": None, "gush_helka": None,
        "start_date": None, "end_date": "2027-01-31", "amounts": None, "pages": None}


def skip(reason):
    try:
        import pytest
        pytest.skip(reason)
    except ImportError:
        print("SKIP:", reason)


def _agreement(conn, file_id, name, metadata=None, version=None, end_date=None):
    conn.execute(
        """insert into agreements (file_id, name, url, modified_time, method, summary, text, embedding,
                                   metadata, metadata_version, end_date)
           values (%s, %s, %s, '1', 'text', 'תקציר קיים', %s, %s::vector, %s, %s, %s)""",
        (file_id, name, f"https://drive/{file_id}", f"טקסט של {name}", VEC,
         None if metadata is None else store.Jsonb(metadata), version, end_date),
    )


@contextlib.contextmanager
def _sync_fakes(conn, listing, downloads):
    """sync.py מול ה-Postgres של הבדיקה, עם Drive, OCR ומודל מזויפים."""
    def fake_extract(f, data):
        if f["id"] == "bad":
            raise RuntimeError("OCR החזיר 3 תווים ל-1 עמודים")
        text = "[עמוד 1]\n\nטקסט\n\n[עמוד 2]\n\nהמשך"  # כמו extract האמיתי במסמך של כמה עמודים
        return (text, "ocr", 91) if f["mimeType"] == PNG else (text, "text", None)

    env = {k: v for k, v in os.environ.items() if k != "GITHUB_STEP_SUMMARY"} | {"DRIVE_FOLDER_ID": "folder"}
    with mock.patch.multiple(store, connect=lambda: contextlib.nullcontext(conn),
                             analyze=lambda name, text, usage=None: ("תקציר", dict(META)),
                             embed=lambda text, usage=None: VEC), \
            mock.patch.multiple(sync.drive, service=lambda: None, list_files=lambda svc, folder: listing,
                                download=lambda svc, f: downloads.append(f["id"]) or b"data"), \
            mock.patch.multiple(sync.extract, extract=fake_extract, page_count=lambda f, data: 3), \
            mock.patch.dict(os.environ, env, clear=True):
        yield


def test_sync_files_table_failed_unsupported_retry():
    with postgres() as conn:
        if conn is None:
            return skip("אין py-pglite או Node")
        store.init(conn)
        for fid, name in (("same", "ללא שינוי.pdf"), ("old", "ישן.pdf"), ("gone", "נמחק מ-Drive.pdf")):
            _agreement(conn, fid, name, metadata=META, version=store.METADATA_VERSION)
            conn.execute("insert into files (file_id, name, status) values (%s, %s, 'indexed')", (fid, name))
        conn.execute("update files set retry_requested = true, status = 'queued' where file_id = 'same'")
        listing = [
            {"id": "new", "name": "סריקה.png", "mimeType": PNG, "modifiedTime": "1", "size": "2048"},
            {"id": "bad", "name": "פגום.pdf", "mimeType": PDF, "modifiedTime": "1", "size": "10"},
            {"id": "video", "name": "סרטון.mp4", "mimeType": "video/mp4", "modifiedTime": "1"},
            {"id": "same", "name": "ללא שינוי.pdf", "mimeType": PDF, "modifiedTime": "1"},
            {"id": "old", "name": "ישן.pdf", "mimeType": PDF, "modifiedTime": "2"},
        ]
        downloads = []
        with _sync_fakes(conn, listing, downloads), mock.patch.object(sync, "MAX_REMOVE", 0.5):  # 1 מתוך 3, מכוון
            stats = sync.sync()

        rows = {r["file_id"]: r for r in store._dicts(conn, "select * from files")}
        assert set(rows) == {"new", "bad", "video", "same", "old"}, rows.keys()  # מה שנמחק מ-Drive נמחק מהטבלה
        assert rows["new"]["status"] == "indexed" and rows["new"]["size_bytes"] == 2048 and rows["new"]["last_attempt_at"]
        assert rows["bad"]["status"] == "failed" and rows["bad"]["error"].startswith("OCR החזיר")
        assert rows["video"]["status"] == "unsupported" and rows["video"]["error"] == "סוג קובץ לא נתמך: video/mp4"
        assert rows["same"]["status"] == "indexed" and rows["same"]["retry_requested"] is False
        assert rows["same"]["last_attempt_at"] is not None
        assert sorted(downloads) == ["bad", "new", "old", "same"]  # same עובד מחדש בגלל retry; סרטון (לא נתמך) לא הורד
        assert [e["file"] for e in stats["errors"]] == ["פגום.pdf"]  # "לא נתמך" אינו שגיאה
        assert stats["unsupported"] == 1 and stats["files"]["added"] == ["סריקה.png"]
        meta = conn.execute("select metadata->>'pages', end_date, ocr_quality from agreements where file_id = 'new'").fetchone()
        assert meta == ("3", date(2027, 1, 31), 91), meta  # מספר העמודים האמיתי, וציון ה-OCR נשמר
        assert conn.execute("select ocr_quality from agreements where file_id = 'old'").fetchone() == (None,)

        downloads.clear()
        with _sync_fakes(conn, listing, downloads):
            sync.sync()
        assert downloads == ["bad"], downloads  # בריצה הבאה רק מה שנכשל מנוסה שוב; ה-retry כבר התאפס


def test_backfill_selects_only_missing_or_stale():
    with postgres() as conn:
        if conn is None:
            return skip("אין py-pglite או Node")
        store.init(conn)
        _agreement(conn, "none", "בלי נתונים.pdf")
        _agreement(conn, "done", "הושלם.pdf", metadata=META, version=store.METADATA_VERSION)
        _agreement(conn, "stale", "גרסה ישנה.pdf", metadata=META, version=store.METADATA_VERSION - 1)
        _agreement(conn, "fail", "ייכשל.pdf")
        assert {d["file_id"] for d in store.pending_metadata(conn)} == {"none", "stale", "fail"}

        def fake_analyze(name, text, usage=None):
            assert text == f"טקסט של {name}"  # מהמסד, בלי Drive ובלי OCR
            if name == "ייכשל.pdf":
                raise RuntimeError("timeout")
            return "תקציר חדש", dict(META)

        stats = {"errors": [], "files": {"backfilled": []}}
        with mock.patch.object(store, "analyze", fake_analyze):
            sync.backfill(conn, stats)
        assert sorted(stats["files"]["backfilled"]) == ["בלי נתונים.pdf", "גרסה ישנה.pdf"]
        assert stats["errors"] == [{"file": "ייכשל.pdf", "error": "השלמת נתונים מובנים: timeout"}]
        assert {d["file_id"] for d in store.pending_metadata(conn)} == {"fail"}  # ינוסה שוב בריצה הבאה
        row = conn.execute("select summary, end_date, metadata_version from agreements where file_id = 'none'").fetchone()
        assert row == ("תקציר קיים", date(2027, 1, 31), store.METADATA_VERSION)  # התקציר וה-embedding לא משתנים


def test_files_panel_queries():
    with postgres() as conn:
        if conn is None:
            return skip("אין py-pglite או Node")
        store.init(conn)
        _agreement(conn, "doc", "הסכם.pdf", META | {"plan_number": 'תת"ל 120'}, 1)
        for fid, status, error in (("doc", "indexed", None), ("broken", "failed", "OCR נכשל")):
            conn.execute("insert into files (file_id, name, status, error) values (%s, %s, %s, %s)",
                         (fid, fid + ".pdf", status, error))
        store.add_queued(conn, {"id": "up", "name": "הועלה.pdf", "mimeType": PDF, "size": "99", "modifiedTime": "t"})

        listed = {f["file_id"]: f for f in store.files_list(conn)}
        assert listed["up"]["url"] == "https://drive.google.com/file/d/up/view" and listed["up"]["status"] == "queued"
        assert listed["doc"]["metadata"]["plan_number"] == 'תת"ל 120' and listed["doc"]["url"] == "https://drive/doc"
        assert listed["doc"]["method"] == "text" and listed["up"]["method"] is None  # לחלון הפרטים: OCR או ישירות
        assert listed["broken"]["error"] == "OCR נכשל"
        assert store.mark_retry(conn, "broken") == "broken.pdf" and store.mark_retry(conn, "missing") is None
        assert store.retry_ids(conn) == {"broken"}

        with mock.patch.object(store, "connect", lambda: contextlib.nullcontext(conn)):
            files = TestClient(main.app).get("/api/files").json()
        assert [f["name"] for f in files] == ["broken.pdf", "doc.pdf", "הועלה.pdf"]


def test_history_save_list_open_delete_rate_and_no_expiry():
    with postgres() as conn:
        if conn is None:
            return skip("אין py-pglite או Node")
        store.init(conn)
        msgs = [{"role": "user", "content": "הסכם שכירות"}, {"role": "assistant", "content": "ט\n\nגוף", "isDoc": True}]
        body = {"title": "הסכם שכירות", "html": "<p><b>גוף</b></p>", "text": "גוף", "messages": msgs,
                "sources": [{"name": "א.pdf", "base": True}]}
        with mock.patch.object(store, "connect", lambda: contextlib.nullcontext(conn)):
            c = TestClient(main.app)
            first = c.post("/api/history", json=body).json()["id"]
            assert c.post("/api/history", json=body | {"id": first, "text": "גוף ערוך"}).json()["id"] == first
            second = c.post("/api/history", json=body | {"title": "שני"}).json()["id"]
            listed = c.get("/api/history").json()
            assert [h["id"] for h in listed] == [second, first]  # האחרון שעודכן ראשון
            assert listed[1]["preview"] == "גוף ערוך" and "html" not in listed[1]

            got = c.get(f"/api/history/{first}").json()
            assert got["html"] == "<p><b>גוף</b></p>" and got["messages"] == msgs and got["sources"][0]["base"]

            # בלי הגבלת זמן: טיוטה ישנה נשארת גם אחרי שמירה של אחרת
            conn.execute("update drafts set updated_at = now() - interval '400 days' where id = %s", (first,))
            c.post("/api/history", json=body | {"id": second})
            assert [h["id"] for h in c.get("/api/history").json()] == [second, first]

            # משוב: נשמר, אפשר לשנות, ומוחזר עם הטיוטה
            assert c.put(f"/api/history/{first}/rating", json={"rating": 1}).json() == {"ok": True}
            assert c.put(f"/api/history/{first}/rating", json={"rating": -1}).status_code == 200
            assert c.get(f"/api/history/{first}").json()["rating"] == -1
            assert c.put(f"/api/history/{first}/rating", json={"rating": 2}).status_code == 422
            assert c.put("/api/history/999999/rating", json={"rating": 1}).status_code == 404

            assert c.delete(f"/api/history/{second}").json() == {"ok": True}
            assert c.delete(f"/api/history/{second}").status_code == 404
            assert c.get(f"/api/history/{second}").status_code == 404
            assert c.post("/api/history", json=body | {"html": "x" * 1_000_001}).status_code == 422


def test_search_learns_from_draft_ratings():
    with postgres() as conn:
        if conn is None:
            return skip("אין py-pglite או Node")
        store.init(conn)
        _agreement(conn, "near", "קרוב.pdf")
        _agreement(conn, "liked", "אהוב.pdf")
        far = "[" + ",".join(["0.1"] * 1535 + ["0.2"]) + "]"  # קצת רחוק יותר מהשאילתה
        conn.execute("update agreements set embedding = %s::vector where file_id = 'liked'", (far,))
        for _ in range(3):
            d = store.save_draft(conn, None, "ט", "<p>x</p>", "x", [], [{"name": "אהוב.pdf"}])
            store.rate_draft(conn, d, 1)
        with mock.patch.object(store, "embed", lambda text, usage=None: VEC):
            assert [r["name"] for r in store.search(conn, "q", k=2)] == ["קרוב.pdf", "אהוב.pdf"]
            assert [r["name"] for r in store.search(conn, "q", k=2, learn=True)] == ["אהוב.pdf", "קרוב.pdf"]


def test_costs_per_draft_and_week_total():
    assert store.cost("gpt-4o", 1_000_000, 400_000, 100_000) == 0.6 * 2.5 + 0.4 * 1.25 + 0.1 * 10
    run = SimpleNamespace(metrics=SimpleNamespace(input_tokens=10_000, cache_read_tokens=0, output_tokens=2_000))
    assert store.agent_usage(run, "gpt-4o") == ("gpt-4o", 10_000, 0, 2_000)
    assert store.agent_usage(SimpleNamespace(), "gpt-4o") is None  # ריצה בלי metrics לא נרשמת
    with postgres() as conn:
        if conn is None:
            return skip("אין py-pglite או Node")
        store.init(conn)
        draft_run = store.agent_usage(run, "gpt-4o")  # 0.025 + 0.02 = $0.045
        store.log_usage(conn, "draft", [("text-embedding-3-small", 50, 0, 0), draft_run, None], "conv-a")
        store.log_usage(conn, "draft", [draft_run], "conv-a")  # תיקון באותה שיחה
        store.log_usage(conn, "draft", [draft_run], "conv-questions-only")
        store.log_usage(conn, "sync", [("gpt-4o", 100_000, 0, 1_000)])
        store.log_usage(conn, "sync", [("gpt-4o", 1_000_000, 0, 0)])
        conn.execute("update api_usage set at = now() - interval '8 days' where input_tokens = 1000000")  # מחוץ לשבוע
        store.save_draft(conn, None, "טיוטה א", "<p>x</p>", "x", [], [], "conv-a")
        store.save_draft(conn, None, "טיוטה ישנה", "<p>y</p>", "y", [], [])  # מלפני מדידת העלויות

        with mock.patch.object(store, "connect", lambda: contextlib.nullcontext(conn)):
            c = TestClient(main.app).get("/api/costs").json()
        assert c["days"] == 7
        with mock.patch.object(store, "connect", lambda: contextlib.nullcontext(conn)):
            c30 = TestClient(main.app).get("/api/costs?days=30").json()
            assert TestClient(main.app).get("/api/costs?days=8").status_code == 400
            assert TestClient(main.app).get("/api/costs/xlsx?days=90").status_code == 200
        assert c30["days"] == 30 and round(c30["sync"]["cost"] - c["sync"]["cost"], 6) == 2.5  # הסנכרון מלפני 8 ימים
        rows = {d["title"]: d for d in c["drafts"]}
        assert round(rows["טיוטה א"]["cost"], 6) == round(2 * 0.045 + 50 * 0.02 / 1e6, 6)
        assert rows["טיוטה א"]["runs"] == 2  # שני ניסוחים; ה-embedding של החיפוש לא נספר כבקשה
        assert rows["טיוטה ישנה"]["cost"] is None
        assert round(c["unsaved"], 6) == 0.045
        assert round(c["draft"]["cost"], 6) == round(3 * 0.045 + 50 * 0.02 / 1e6, 6) and c["draft"]["calls"] == 4
        assert round(c["sync"]["cost"], 6) == 0.26 and c["sync"]["calls"] == 1
        assert round(c["total"], 6) == round(c["draft"]["cost"] + c["sync"]["cost"], 6)

        # ניסוח באתר רושם את העלות עם מזהה השיחה, גם בלי מסמכי בסיס
        def fake_run(self, messages):
            return SimpleNamespace(status=RunStatus.completed, metrics=run.metrics,
                                   content=draft.Draft(status="questions", questions=["מי הצדדים?"], title="", body=""))
        with mock.patch.object(store, "connect", lambda: contextlib.nullcontext(conn)), \
                mock.patch.object(Agent, "run", fake_run):
            assert TestClient(main.app).post("/api/draft", json={"messages": [{"role": "user", "content": "נסח"}],
                                                                  "conversation": "conv-b"}).status_code == 200
        assert conn.execute("select count(*) from api_usage where conversation = 'conv-b'").fetchone()[0] == 1


def test_ask_cites_only_documents_the_tools_returned():
    from app import ask
    doc = {"file_id": "d1", "name": "הסכם.pdf", "url": "https://drive/d1", "summary": "s",
           "text": "[עמוד 1]\n\nמבוא\n\n[עמוד 2]\n\n7.2 הסודיות תחול 5 שנים"}
    logged = []

    def fake_run(self, messages):
        self.tools[0]("סודיות")  # search_documents: המודל "ראה" את d1
        assert "[עמוד 2]" in self.tools[2]("d1") and self.tools[2]("nope") == "אין מסמך עם המזהה הזה."
        answer = ask.Answer(answer="הסודיות היא 5 שנים [1]. וגם [2]", citations=[
            ask.Citation(n=1, file_id="d1", section="סעיף 7.2", page=2, quote="הסודיות תחול 5 שנים"),
            ask.Citation(n=2, file_id="invented", section="", page=None, quote="x")])
        return SimpleNamespace(status=RunStatus.completed, content=answer, metrics=SimpleNamespace(
            input_tokens=1000, cache_read_tokens=0, output_tokens=100))

    with mock.patch.multiple(store, connect=lambda: contextlib.nullcontext(None),
                             search=lambda conn, q, k=3, usage=None, learn=False: [doc], get_doc=lambda conn, fid: doc if fid == "d1" else None,
                             record=lambda kind, usage, conv=None: logged.append((kind, usage, conv))), \
            mock.patch.object(Agent, "run", fake_run):
        r = TestClient(main.app).post("/api/ask", json={"messages": [{"role": "user", "content": "כמה זמן הסודיות?"}],
                                                        "conversation": "q1"}).json()
    assert r["citations"] == [{"n": 1, "file_id": "d1", "section": "סעיף 7.2", "page": 2,
                               "quote": "הסודיות תחול 5 שנים", "name": "הסכם.pdf", "url": "https://drive/d1"}]
    assert logged == [("ask", [("gpt-4o", 1000, 0, 100)], "q1")]


def test_draft_web_check_and_deep_review():
    calls, logged = [], []
    reviewed = {"ok": True}

    def fake_run(self, messages):
        model = self.model.id
        calls.append((model, [t for t in self.tools if isinstance(t, dict)] if self.tools else []))
        m = SimpleNamespace(input_tokens=100, cache_read_tokens=0, output_tokens=10)
        if model == draft.REVIEW_MODEL:
            if not reviewed["ok"]:
                raise TimeoutError("timeout")
            assert "<טיוטה>" in messages and "נטיית הניסוח" in messages
            return SimpleNamespace(status=RunStatus.completed, metrics=m,
                                   content=draft.Review(title="הסכם", body="גוף מתוקן", notes=["הוסף סעיף שיפוי"]))
        assert ("דרישות חוק" in (self.additional_context or "")) == bool(calls[-1][1])
        return SimpleNamespace(status=RunStatus.completed, metrics=m,
                               content=draft.Draft(status="agreement", questions=[], title="הסכם", body="גוף"))

    msgs = [{"role": "user", "content": "סוג החוזה: הסכם שירותים\n\nנטיית הניסוח: מאוזנת"}]
    with mock.patch.object(Agent, "run", fake_run), \
            mock.patch.object(store, "record", lambda kind, usage, conv=None: logged.append(usage)):
        out = draft.draft(msgs, web=True, review=True)
        assert out["body"] == "גוף מתוקן" and out["review_notes"] == ["הוסף סעיף שיפוי"]
        assert calls == [("gpt-4o", [{"type": "web_search"}]), ("gpt-5.5", [])]
        assert [u[0] for u in logged[-1]] == ["gpt-4o", "web_search", "gpt-5.5"]
        assert round(store.cost("web_search", 1), 4) == 0.025

        reviewed["ok"] = False  # הבדיקה נכשלה: הטיוטה המקורית חוזרת עם הערה, והעלות של הניסוח עדיין נרשמת
        out = draft.draft(msgs, review=True)
        assert out["body"] == "גוף" and "נכשלה" in out["review_notes"][0]
        assert [u[0] for u in logged[-1]] == ["gpt-4o"]

        calls.clear()
        draft.draft(msgs)  # בלי האפשרויות: בלי חיפוש באינטרנט ובלי בדיקה
        assert calls == [("gpt-4o", [])]


def test_server_start_creates_tables_before_first_sync():
    with postgres() as conn:
        if conn is None:
            return skip("אין py-pglite או Node")
        with mock.patch.object(store, "connect", lambda: contextlib.nullcontext(conn)):
            with TestClient(main.app) as c:  # with מריץ את ה-lifespan, כמו הפעלת השרת ב-Render
                for _ in range(100):  # המיגרציות רצות ברקע, אחרי שהשרת כבר עונה
                    if conn.execute("select to_regclass('files')").fetchone()[0]:
                        break
                    time.sleep(0.1)
                r = c.get("/api/files")
                assert r.status_code == 200 and r.json() == [], r.text  # "הרשימה תתמלא אחרי הסנכרון הבא"
                assert c.get("/api/reports/index").status_code == 404  # הדוחות הוסרו


def test_server_answers_before_slow_migrations_and_constraint_upgrades_once():
    # מסד שלא עונה: השרת עונה מיד, והמיגרציות נכשלות ברקע (ב-Render: "Port scan timeout")
    started = threading.Event()

    def hang():
        started.set()
        time.sleep(3)
        raise OSError("connection timeout expired")

    with mock.patch.object(store, "connect", hang), mock.patch.dict(os.environ, {"RENDER": ""}):
        t0 = time.monotonic()
        with TestClient(main.app) as c:
            assert c.get("/healthz").json() == {"ok": True} and time.monotonic() - t0 < 2
        assert started.is_set()

    with postgres() as conn:
        if conn is None:
            return skip("אין py-pglite או Node")
        store.init(conn)
        # מסד מלפני 'ask': ה-constraint הישן מוחלף פעם אחת, ואחר כך init לא נוגע בו
        conn.execute("alter table api_usage drop constraint api_usage_kind_check")
        conn.execute("alter table api_usage add constraint api_usage_kind_check check (kind in ('draft', 'sync'))")
        store.init(conn)
        store.log_usage(conn, "ask", [("gpt-4o", 1, 0, 1)])
        oid = conn.execute("select oid from pg_constraint where conname = 'api_usage_kind_check'").fetchone()[0]
        store.init(conn)
        assert conn.execute("select oid from pg_constraint where conname = 'api_usage_kind_check'").fetchone()[0] == oid
        assert conn.execute("show lock_timeout").fetchone()[0] == "0"  # רק למיגרציות, לא לשאר החיבור


def _docx_bytes():
    import docx
    buf = io.BytesIO()
    docx.Document().save(buf)
    return buf.getvalue()


def _heic_bytes(img=None):
    """תמונת HEIC אמיתית, כמו שאייפון שומר (pillow-heif כותב HEIC)."""
    from PIL import Image
    buf = io.BytesIO()
    (img or Image.new("RGB", (64, 64), "white")).save(buf, "HEIF")
    return buf.getvalue()


def test_old_word_doc_is_converted_through_google_docs():
    """DOC (Word 97-2003): עותק זמני מומר ל-Google Docs, מייצאים טקסט, והעותק נמחק גם אם הייצוא נכשל."""
    calls = []

    class Req:
        def __init__(self, value):
            self.value = value

        def execute(self):
            return self.value

    class FakeDrive:
        def files(self):
            return self

        def copy(self, fileId, body, **kw):
            calls.append(("copy", fileId, body["mimeType"], body["parents"]))
            return Req({"id": "tmp1"})

        def export_media(self, fileId, mimeType):
            calls.append(("export", fileId, mimeType))
            return "export-request"

        def delete(self, fileId, **kw):
            calls.append(("delete", fileId))
            return Req(None)

    f = {"id": "doc1", "name": "הסכם ישן.doc", "mimeType": extract.DOC}
    with mock.patch.object(drive, "_fetch", lambda req: "הסכם שכירות ישן".encode()):
        data = drive.download(FakeDrive(), f)
    assert calls == [("copy", "doc1", drive.GDOC, ["root"]), ("export", "tmp1", "text/plain"), ("delete", "tmp1")]
    assert extract.extract(f, data) == ("הסכם שכירות ישן", "text", None)

    calls.clear()

    def failing(req):
        raise RuntimeError("export failed")

    with mock.patch.object(drive, "_fetch", failing):
        try:
            drive.download(FakeDrive(), f)
            raise AssertionError("היה צריך להיכשל")
        except RuntimeError:
            pass
    assert calls[-1] == ("delete", "tmp1")  # העותק הזמני לא נשאר ב-Drive


def test_upload_validation():
    ok = {"a.pdf": b"%PDF-1.7", "a.PNG": b"\x89PNG\r\n\x1a\n...", "a.jpg": b"\xff\xd8\xff\xe0", "a.jpeg": b"\xff\xd8\xff",
          "a.tif": b"II*\x00..", "a.tiff": b"MM\x00*..", "a.webp": b"RIFF\x00\x00\x00\x00WEBPVP8 ", "a.docx": _docx_bytes(),
          "a.doc": b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1....", "IMG_1234.HEIC": _heic_bytes(), "a.heif": _heic_bytes()}
    for name, data in ok.items():
        assert drive.check_upload(name, data)[0] == name
    assert drive.check_upload("C:\\Users\\x\\חוזה.pdf", b"%PDF-")[0] == "חוזה.pdf"  # בלי נתיב
    zip_not_docx = io.BytesIO()
    zipfile.ZipFile(zip_not_docx, "w").writestr("x.txt", "x")
    bad = {"a.exe": b"MZ", "a": b"%PDF-", "a.pdf": b"\x89PNG\r\n\x1a\n", "a.png": b"%PDF-", "b.pdf": b"",
           "big.pdf": b"%PDF-" + b"0" * drive.MAX_UPLOAD, "a.docx": zip_not_docx.getvalue(),
           "a.heic": b"\x89PNG\r\n\x1a\n...", "a.doc": b"%PDF-1.7"}
    for name, data in bad.items():
        try:
            drive.check_upload(name, data)
            raise AssertionError(f"היה צריך לדחות: {name}")
        except ValueError:
            pass
    assert drive.unique_name("הסכם.pdf", set()) == "הסכם.pdf"
    assert drive.unique_name("הסכם.pdf", {"הסכם.pdf"}) == "הסכם (1).pdf"
    assert drive.unique_name("הסכם.pdf", {"הסכם.pdf", "הסכם (1).pdf"}) == "הסכם (2).pdf"


class UploadDrive:
    """Drive מזויף להעלאה: בתיקייה כבר יש "הסכם.pdf", ו-get מחזיר את מה שנשמר, עם השינויים שב-wrong."""
    def __init__(self):
        self.wrong, self.saved = {}, None

    def files(self):
        return self

    def list(self, **kw):
        return SimpleNamespace(execute=lambda: {"files": [{"name": "הסכם.pdf"}]})

    def create(self, body, media_body, **kw):
        data = media_body.getbytes(0, media_body.size())
        self.saved = {"parents": body["parents"], "size": str(len(data)), "md5Checksum": hashlib.md5(data).hexdigest(),
                      "webViewLink": "https://drive.google.com/file/d/x/view"}
        return SimpleNamespace(execute=lambda: {"id": "x", "name": body["name"], "mimeType": PDF})

    def get(self, fileId, fields, **kw):
        assert fileId == "x" and "md5Checksum" in fields and "parents" in fields
        return SimpleNamespace(execute=lambda: self.saved | self.wrong)


def test_upload_api_requires_password_and_renames():
    calls = []
    with mock.patch.multiple(main, PASSWORD="", TOKEN="t"), mock.patch.dict(os.environ, {"RENDER": ""}), \
            mock.patch.object(drive, "service", lambda: calls.append("drive")):
        c = TestClient(main.app)  # פיתוח מקומי בלי סיסמה: האתר פתוח, אבל פעולות כתיבה חסומות
        assert c.get("/api/me").json() == {"writer": False, "sync": github.configured()}
        assert c.post("/api/upload?name=הסכם.pdf", content=b"%PDF-1.7").status_code == 403
        assert c.post("/api/sync").status_code == 403
    assert calls == []  # לא מחובר: אין קריאה ל-Drive

    queued, fake = [], UploadDrive()
    with mock.patch.multiple(main, PASSWORD="secret", TOKEN="token-of-secret"), \
            mock.patch.object(drive, "service", lambda: fake), \
            mock.patch.dict(os.environ, {"DRIVE_FOLDER_ID": "folder"}), \
            mock.patch.multiple(store, connect=lambda: contextlib.nullcontext(None),
                                add_queued=lambda conn, f: queued.append(f["name"])), \
            mock.patch.object(github, "configured", lambda: False):
        c = TestClient(main.app)
        assert c.post("/api/upload?name=הסכם.pdf", content=b"%PDF-1.7").status_code == 401  # לא מחובר
        assert fake.saved is None
        assert c.post("/login", data={"password": "secret"}, follow_redirects=False).status_code == 303
        assert c.get("/api/me").json()["writer"] is True
        r = c.post("/api/upload?name=הסכם.pdf", content=b"%PDF-1.7")
        assert r.json() == {"name": "הסכם (1).pdf", "renamed": True, "sync": None,
                            "webViewLink": "https://drive.google.com/file/d/x/view"}, r.text
        assert queued == ["הסכם (1).pdf"]

        # מה ש-Drive שמר לא תואם למה שנשלח: שגיאה, ואין שורה בטבלת files
        for wrong in ({"parents": ["other"]}, {"size": "1"}, {"md5Checksum": "0" * 32}):
            fake.wrong = wrong
            r = c.post("/api/upload?name=הסכם.pdf", content=b"%PDF-1.7")
            assert r.status_code == 502 and r.json()["detail"] == "הקובץ לא נשמר בתיקייה כמו שצריך", (wrong, r.text)
        assert queued == ["הסכם (1).pdf"]
        assert c.post("/api/upload?name=a.pdf", content=b"\x89PNG\r\n\x1a\n").status_code == 400  # תוכן לא תואם
        big = c.post("/api/upload?name=a.pdf", content=b"%PDF-" + b"0" * drive.MAX_UPLOAD)
        assert big.status_code == 413


def test_draft_limits_and_errors():
    ok = {"status": "questions", "questions": ["?"], "title": "", "body": "", "sources": []}
    body = {"messages": [{"role": "user", "content": "נסח"}]}
    with mock.patch.multiple(main, PASSWORD="", TOKEN="t", DAILY_DRAFT_LIMIT=40, _daily={"date": None, "count": 0},
                             _drafts=main.defaultdict(main.deque)), \
            mock.patch.dict(os.environ, {"RENDER": ""}), mock.patch.object(draft, "draft", lambda *a: ok):
        a, b = TestClient(main.app, cookies={"sid": "a"}), TestClient(main.app, cookies={"sid": "b"})
        assert all(a.post("/api/draft", json=body).status_code == 200 for _ in range(main.DRAFTS_PER_HOUR))
        r = a.post("/api/draft", json=body)
        assert r.status_code == 429 and "ניסוחים בשעה" in r.json()["detail"]
        assert all(b.post("/api/draft", json=body).status_code == 200 for _ in range(10))  # דפדפן אחר: מגבלה נפרדת
        r = b.post("/api/draft", json=body)  # 40 ניסוחים באתר היום
        assert r.status_code == 429 and "היומית" in r.json()["detail"]

    def boom(m, f):
        raise RuntimeError("openai 500: secret-internal-detail")

    with mock.patch.multiple(main, PASSWORD="", TOKEN="t", _daily={"date": None, "count": 0}), \
            mock.patch.dict(os.environ, {"RENDER": ""}), mock.patch.object(draft, "draft", boom):
        r = TestClient(main.app).post("/api/draft", json=body)
    assert r.status_code == 502 and "הניסוח נכשל" in r.json()["detail"] and "secret" not in r.text
    model = draft.build_agent([]).model
    assert (model.timeout, model.max_retries) == (120, 1)


def test_whole_site_requires_login():
    with mock.patch.multiple(main, PASSWORD="secret", TOKEN="token-of-secret"):
        c = TestClient(main.app, follow_redirects=False)
        r = c.get("/")
        assert r.status_code == 303 and r.headers["location"] == "/login?next=/"  # מסך כניסה בפעם הראשונה
        for path in ("/api/files", "/api/status", "/api/status/docx", "/admin/sync"):
            assert c.get(path).status_code == 401, path
        assert c.post("/api/draft", json={"messages": [{"role": "user", "content": "x"}]}).status_code == 401
        assert c.get("/healthz").status_code == 200  # בדיקת התקינות של Render
        assert c.get("/static/logo.webp").status_code == 200 and c.get("/static/favicon.png").status_code == 200
        assert "סיסמה" in c.get("/login").text

        wrong = c.post("/login", data={"password": "nope", "next": "/admin/sync"})
        assert wrong.status_code == 303 and wrong.headers["location"] == "/login?error=1&next=/admin/sync"
        assert "session" not in c.cookies
        for evil in ("//evil.com", "https://evil.com", "/\\evil.com"):  # לא מפנים לאתר אחר אחרי הכניסה
            assert c.post("/login", data={"password": "secret", "next": evil}).headers["location"] == "/"
        ok = c.post("/login", data={"password": "secret", "next": "/admin/sync"})
        assert ok.status_code == 303 and ok.headers["location"] == "/admin/sync"
        cookie = ok.headers["set-cookie"].lower()
        assert "max-age=2592000" in cookie and "httponly" in cookie  # 30 יום
        assert c.get("/").status_code == 200

    with mock.patch.multiple(main, PASSWORD="", TOKEN="t"):
        with mock.patch.dict(os.environ, {"RENDER": "true"}):  # ב-Render בלי סיסמה: חסום, לא פתוח לכולם
            c = TestClient(main.app, follow_redirects=False)
            assert c.get("/").status_code == 503 and c.get("/healthz").status_code == 200
        with mock.patch.dict(os.environ, {"RENDER": ""}):  # פיתוח מקומי: פתוח
            assert TestClient(main.app).get("/").status_code == 200


def test_focus_file_in_agent_context():
    base = {"file_id": "b1", "name": "בסיס.pdf", "url": "https://drive/b1", "summary": "תקציר", "text": "טקסט הבסיס"}
    sources = []
    agent = draft.build_agent(sources, [base])
    assert "טקסט הבסיס" in agent.additional_context and "הדוגמאות העיקריות" in agent.additional_context
    assert sources == [{"name": "בסיס.pdf", "url": "https://drive/b1", "summary": "תקציר", "base": True}]
    found = [base, {"name": "אחר.pdf", "url": "u", "summary": "s", "text": "t"}]
    with mock.patch.multiple(store, connect=lambda: contextlib.nullcontext(None), search=lambda conn, q, usage=None, learn=False: found):
        agent.tools[0]("שאילתה")
    assert [s["name"] for s in sources] == ["בסיס.pdf", "אחר.pdf"]  # הבסיס לא נכפל, ונשאר ראשון

    def fake_run(self, messages):
        assert "טקסט הבסיס" in self.additional_context
        return SimpleNamespace(status=RunStatus.completed,
                               content=draft.Draft(status="agreement", questions=[], title="הסכם", body="גוף"))

    # כמה מסמכי בסיס: כולם בהקשר ובמקורות, בסדר הבחירה ובלי כפילויות; מעבר ל-3 כל אחד מקוצר
    docs = {"b1": base, "b2": {**base, "file_id": "b2", "name": "בסיס2.pdf", "text": "טקסט שני"}}
    with mock.patch.multiple(store, connect=lambda: contextlib.nullcontext(None),
                             get_doc=lambda conn, fid: docs.get(fid)), mock.patch.object(Agent, "run", fake_run):
        out = draft.draft([{"role": "user", "content": "נסח"}], ["b2", "b1", "b2"])
        assert [(s["name"], s.get("base")) for s in out["sources"]] == [("בסיס2.pdf", True), ("בסיס.pdf", True)]
        assert out["title"] == "הסכם"
        r = TestClient(main.app).post("/api/draft", json={"messages": [{"role": "user", "content": "נסח"}],
                                                          "focus_file_ids": ["b1", "missing"]})
        assert r.status_code == 404 and "לא נמצא" in r.json()["detail"]
    many = [{**base, "name": f"{i}.pdf", "text": "x" * draft.EXAMPLE_CHARS} for i in range(6)]
    ctx = draft.build_agent([], many).additional_context
    assert ctx.count("<מסמך בסיס") == 6 and ctx.count("x") == 6 * (draft.BASE_CHARS // 6)


def test_sync_trigger_rate_limit():
    calls = []
    with mock.patch.object(github, "_api", lambda method, path, body=None: calls.append(path)), \
            mock.patch.multiple(github, _last=None, _timer=None):
        assert github.trigger()["started"] is True
        second = github.trigger()
        assert second["started"] is False and 290 < second["retry_after"] <= 301
        deferred = github.trigger(defer=True)
        assert deferred["deferred"] is True and github._timer is not None
        github._timer.cancel()
        assert calls == ["actions/workflows/sync.yml/dispatches"]  # רק הפעלה אחת בפועל


def test_image_ocr_with_tesseract():
    if not (shutil.which("tesseract") or Path(extract.WINDOWS_TESSERACT).exists()):
        return skip("Tesseract לא מותקן")
    from PIL import Image, ImageDraw, ImageFont
    font = next((ImageFont.truetype(f, 36) for f in ("arial.ttf", "DejaVuSans.ttf")
                 if _font_exists(ImageFont, f)), None)
    if font is None:
        return skip("אין גופן TrueType לציור טקסט")

    def page(text):  # לפחות 50 תווים לעמוד, אחרת extract מסמן את ה-OCR ככושל (בכוונה)
        img = Image.new("RGB", (1400, 360), "white")
        ImageDraw.Draw(img).multiline_text((40, 40), text, fill="black", font=font, spacing=20)
        return img

    p1 = page("This agreement is made between\nthe Ministry and the Company\nfor the night parking project.")
    p2 = page("The term of this agreement ends\nin December 2030, unless both\nparties agree otherwise in writing.")
    png, tif = io.BytesIO(), io.BytesIO()
    p1.save(png, "PNG")
    p1.save(tif, "TIFF", save_all=True, append_images=[p2])
    text, method, quality = extract.extract({"mimeType": PNG}, png.getvalue())
    assert method == "ocr" and "agreement" in text, text
    assert 60 <= quality <= 100, quality  # טקסט נקי ומודפס: ביטחון גבוה
    text, method, quality = extract.extract({"mimeType": "image/tiff"}, tif.getvalue())
    assert "agreement" in text and "December" in text, text  # שני העמודים של ה-TIFF
    assert extract.page_count({"mimeType": "image/tiff"}, tif.getvalue()) == 2


def _text_pdf(text):
    """PDF דיגיטלי מינימלי עם שכבת טקסט אמיתית (כמו קובץ שנוצר מ-Word)."""
    stream = f"BT /F1 11 Tf 40 750 Td ({text}) Tj ET".encode()
    objs = [b"<</Type/Catalog/Pages 2 0 R>>", b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
            b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]/Resources<</Font<</F1 4 0 R>>>>/Contents 5 0 R>>",
            b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
            b"<</Length %d>>stream\n" % len(stream) + stream + b"\nendstream"]
    out, offsets = io.BytesIO(b"%PDF-1.4\n"), []
    out.seek(0, 2)
    for i, obj in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % i + obj + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1) + b"".join(b"%010d 00000 n \n" % o for o in offsets))
    out.write(b"trailer<</Size %d/Root 1 0 R>>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref))
    return out.getvalue()


def test_pdf_scanned_vs_digital_detection():
    """PDF עם שכבת טקסט נקרא ישירות, בלי OCR. PDF שהוא רק תמונה (סריקה פיזית) מזוהה ועובר OCR."""
    from PIL import Image
    pdf = {"mimeType": PDF}
    digital = _text_pdf("This agreement is made between the Ministry and the Company for the night parking project.")
    scanned = io.BytesIO()
    Image.new("RGB", (1240, 1754), "white").save(scanned, "PDF")  # עמוד סרוק: תמונה בלבד, אין טקסט לסמן

    ocr_calls = []

    def fake_ocr(images):
        pages = list(images)
        ocr_calls.append(len(pages))
        return ["טקסט שזוהה מהסריקה, ארוך מספיק כדי לעבור את בדיקת האיכות של 50 תווים לעמוד."] * len(pages), 77

    with mock.patch.object(extract, "ocr", fake_ocr):
        text, method, quality = extract.extract(pdf, digital)
        assert method == "text" and "Ministry" in text and ocr_calls == [] and quality is None  # לא נשלח ל-OCR
        text, method, quality = extract.extract(pdf, scanned.getvalue())
        assert method == "ocr" and ocr_calls == [1] and "סריקה" in text and quality == 77  # זוהה כסרוק


def test_iphone_photos_heic_and_rotation():
    """צילום אייפון: HEIC עובר OCR, ותמונה ששמורה על הצד עם תג סיבוב (EXIF) מסובבת לפני ה-OCR."""
    if not (shutil.which("tesseract") or Path(extract.WINDOWS_TESSERACT).exists()):
        return skip("Tesseract לא מותקן")
    from PIL import Image, ImageDraw, ImageFont
    try:
        font = ImageFont.truetype("arial.ttf", 36)
    except OSError:
        return skip("אין גופן TrueType לציור טקסט")
    page = Image.new("RGB", (1400, 360), "white")
    ImageDraw.Draw(page).multiline_text(
        (40, 40), "This agreement is made between\nthe Ministry and the Company\nfor the night parking project.",
        fill="black", font=font, spacing=20)

    text, method, quality = extract.extract({"mimeType": "image/heic"}, _heic_bytes(page))
    assert method == "ocr" and "agreement" in text and quality >= 70, (text, quality)

    sideways = page.rotate(90, expand=True)  # כמו שהמצלמה שומרת את הפיקסלים
    exif = Image.Exif()
    exif[0x0112] = 6  # Orientation: להציג מסובב 90° עם כיוון השעון
    jpg = io.BytesIO()
    sideways.save(jpg, "JPEG", exif=exif, quality=95)
    raw_text, _ = extract.ocr([sideways])
    assert "agreement" not in " ".join(raw_text)  # בלי הסיבוב ה-OCR באמת נכשל, כך שהבדיקה משמעותית
    text, method, quality = extract.extract({"mimeType": "image/jpeg"}, jpg.getvalue())
    assert "agreement" in text and "Ministry" in text, text


def test_docx_with_scanned_pages_goes_through_ocr():
    """Word עם טקסט נקרא ישירות. Word שבתוכו רק עמוד סרוק כתמונה עובר OCR על התמונה ומקבל ציון."""
    import docx
    from PIL import Image
    mime = {"mimeType": extract.DOCX}

    def build(paragraph, with_image):
        d, buf = docx.Document(), io.BytesIO()
        d.add_paragraph(paragraph)
        if with_image:
            img = io.BytesIO()
            Image.new("RGB", (800, 1100), "white").save(img, "PNG")
            img.seek(0)
            d.add_picture(img)
        d.save(buf)
        return buf.getvalue()

    ocr_calls = []

    def fake_ocr(images):
        ocr_calls.append(len(images))
        return ["הסכם שכירות בין המשכיר לבין השוכר, לתקופה של שנתיים, כפי שזוהה מהעמוד הסרוק."], 82

    with mock.patch.object(extract, "ocr", fake_ocr):
        text, method, quality = extract.extract(mime, build("הסכם שכירות. " * 30, with_image=True))
        assert method == "docx" and quality is None and ocr_calls == []  # יש מספיק טקסט: לא מריצים OCR
        text, method, quality = extract.extract(mime, build("נספח א", with_image=True))
        assert method == "ocr" and quality == 82 and ocr_calls == [1], (method, ocr_calls)
        assert text.startswith("נספח א") and "המשכיר" in text  # הטקסט מהקובץ ואחריו הטקסט מהסריקה
        text, method, quality = extract.extract(mime, build("מכתב קצר", with_image=False))
        assert method == "docx" and text == "מכתב קצר"  # טקסט קצר בלי תמונות: כמו שהוא


def _font_exists(image_font, name):
    try:
        image_font.truetype(name, 10)
        return True
    except OSError:
        return False


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    for test in [v for k, v in dict(globals()).items() if k.startswith("test_")]:
        test()
        print("ok ", test.__name__)
